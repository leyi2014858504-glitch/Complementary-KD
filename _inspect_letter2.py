# -*- coding: utf-8 -*-
from docx import Document

d = Document(r"d:\project ML\Response_Letter_Revised.docx")
print("total paragraphs:", len(d.paragraphs))
for i, p in enumerate(d.paragraphs):
    t = p.text
    if "HASKD" in t:
        print(f"HASKD at [{i}]")
    if "\u00a73.1" in t:
        print(f"sec3.1 at [{i}]: {t[:120]}")
print("--- around Comment 5 ---")
for i in range(15, 19):
    print(f"[{i}] runs={len(d.paragraphs[i].runs)}: {d.paragraphs[i].text[:400]}")
